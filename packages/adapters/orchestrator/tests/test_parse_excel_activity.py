"""Tests for ParseExcelActivity — wiring to open_banca_parsing engine."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from temporalio.exceptions import ApplicationError

from open_banca_orchestrator.activities.parse_excel import (
    ParseExcelInput,
    ParserConfig,
    parse_excel,
)


def _make_test_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Ahorro"
    headers = ["Fecha", "Descripcion", "Debito", "Credito", "Saldo", "Referencia"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    data = ["15/03/2024", "Test txn", "0", "100.00", "100.00", "REF"]
    for col, value in enumerate(data, start=1):
        ws.cell(row=2, column=col, value=value)
    wb.save(path)


def test_parse_excel_parses_file(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "statement.xlsx"
    _make_test_xlsx(xlsx_path)

    result = parse_excel(
        ParseExcelInput(
            excel_path=str(xlsx_path),
            content_hash="abc123",
            account_id="acc-fallback",
            parser_config=ParserConfig(bank_id="banco_general"),
        )
    )

    assert result.row_count == 1
    assert len(result.transactions) == 1
    txn = result.transactions[0]
    assert txn.amount == "100.00"
    assert txn.date == "2024-03-15"
    assert txn.currency == "PAB"
    assert txn.fingerprint_hash
    assert txn.description == "Test txn"


def test_parse_excel_missing_file_raises_application_error(tmp_path: Path) -> None:
    with pytest.raises(ApplicationError, match="Excel file not found"):
        parse_excel(
            ParseExcelInput(
                excel_path=str(tmp_path / "missing.xlsx"),
                content_hash="abc",
                account_id="acc-001",
                parser_config=ParserConfig(bank_id="banco_general"),
            )
        )


def test_parse_excel_missing_parser_raises_application_error(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "statement.xlsx"
    _make_test_xlsx(xlsx_path)

    with pytest.raises(ApplicationError, match="parser.json not found"):
        parse_excel(
            ParseExcelInput(
                excel_path=str(xlsx_path),
                content_hash="abc",
                account_id="acc-001",
                parser_config=ParserConfig(bank_id="nonexistent_bank_xyz"),
            )
        )


def test_parse_excel_invalid_xlsx_raises_application_error(tmp_path: Path) -> None:
    bad_path = tmp_path / "bad.xlsx"
    bad_path.write_bytes(b"NOT A VALID XLSX")

    with pytest.raises(ApplicationError):
        parse_excel(
            ParseExcelInput(
                excel_path=str(bad_path),
                content_hash="abc",
                account_id="acc-001",
                parser_config=ParserConfig(bank_id="banco_general"),
            )
        )


def test_transaction_record_has_fingerprint_hash(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "statement.xlsx"
    _make_test_xlsx(xlsx_path)

    result = parse_excel(
        ParseExcelInput(
            excel_path=str(xlsx_path),
            content_hash="abc",
            account_id="acc-001",
            parser_config=ParserConfig(bank_id="banco_general"),
        )
    )
    assert all(len(t.fingerprint_hash) == 64 for t in result.transactions)
