"""Integration tests — banco_general parser.json against in-memory Excel fixtures."""

from __future__ import annotations

import io
import json
from decimal import Decimal
from pathlib import Path

import openpyxl

from open_banca_domain.entities.bank_map import ParserConfig
from open_banca_parsing.engine import ExcelParser

_PARSER_PATH = Path(__file__).parent.parent / "parser.json"


def _load_parser_config() -> ParserConfig:
    data = json.loads(_PARSER_PATH.read_text(encoding="utf-8"))
    return ParserConfig(**data)


def _make_bg_savings_xlsx(rows: list[list[object]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Ahorro"
    headers = ["Fecha", "Descripcion", "Debito", "Credito", "Saldo", "Referencia"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    for row_idx, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parser_json_parses_savings_debit_row() -> None:
    xlsx = _make_bg_savings_xlsx(
        [["15/03/2024", "Retiro ATM", "50.00", "0", "950.00", "REF001"]]
    )
    txns = ExcelParser().parse(xlsx, _load_parser_config())
    assert len(txns) == 1
    assert txns[0].amount == Decimal("-50.00")
    assert txns[0].currency == "PAB"
    assert "ATM" in txns[0].description
    assert txns[0].fingerprint_hash


def test_parser_json_parses_savings_credit_row() -> None:
    xlsx = _make_bg_savings_xlsx(
        [["16/03/2024", "Deposito nómina", "0", "1,500.00", "2,450.00", "REF002"]]
    )
    txns = ExcelParser().parse(xlsx, _load_parser_config())
    assert len(txns) == 1
    assert txns[0].amount == Decimal("1500.00")


def test_parser_json_multiple_rows() -> None:
    xlsx = _make_bg_savings_xlsx(
        [
            ["01/01/2024", "Compra", "25.00", "0", "975.00", "R1"],
            ["02/01/2024", "Abono", "0", "100.00", "1,075.00", "R2"],
        ]
    )
    txns = ExcelParser().parse(xlsx, _load_parser_config())
    assert len(txns) == 2
    assert txns[0].amount == Decimal("-25.00")
    assert txns[1].amount == Decimal("100.00")
